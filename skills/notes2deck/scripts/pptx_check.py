#!/usr/bin/env python3
"""Pre-delivery checks for a presentation .pptx.
Reports: font faces used, notes masters without a theme, question marks on
slides, slides without '（约 N 分钟）' notes, total minutes, banned phrases.
Usage: pptx_check.py DECK.pptx [--ban 词1,词2] [--latin NAME] [--ea NAME]"""
import argparse, collections, re, zipfile
from lxml import etree
from pptx import Presentation
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
THEME_T = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
DEFAULT_BAN = ["汇报总览", "时间安排", "可能会问", "怎么讲", "讲什么", "讲解建议", "博士", "课堂报告"]
ELABEL = re.compile(r"(?<![A-Za-z0-9])E[0-3](?![0-9A-Za-z])")

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("deck")
    ap.add_argument("--ban", default=""); ap.add_argument("--latin", default="Times New Roman"); ap.add_argument("--ea", default="等线")
    a = ap.parse_args(); problems = []
    faces = collections.Counter()
    with zipfile.ZipFile(a.deck) as z:
        for n in z.namelist():
            if n.endswith(".xml") and n.startswith("ppt/"):
                root = etree.fromstring(z.read(n))
                for el in root.iter(f"{{{A}}}latin", f"{{{A}}}ea", f"{{{A}}}cs"):
                    f = el.get("typeface", "")
                    if f and not f.startswith("+"): faces[f] += 1
        master_themes = set()
        for n in z.namelist():
            if n.startswith("ppt/slideMasters/_rels/") and n.endswith(".rels"):
                master_themes |= set(re.findall(r'Type="%s" Target="([^"]+)"' % re.escape(THEME_T), z.read(n).decode()))
        for n in z.namelist():
            if n.startswith("ppt/notesMasters/_rels/") and n.endswith(".rels"):
                tg = re.findall(r'Type="%s" Target="([^"]+)"' % re.escape(THEME_T), z.read(n).decode())
                if not tg or set(tg) & master_themes:
                    problems.append(f"{n}: notes master has no theme of its own (PowerPoint will ask to repair) -> pptx_fonts.py --fix-notes-theme")
    bad = {f: c for f, c in faces.items() if f not in (a.latin, a.ea)}
    if bad: problems.append(f"unexpected fonts: {bad}")
    ban = DEFAULT_BAN + [b for b in a.ban.split(",") if b]
    total = 0
    p = Presentation(a.deck)
    for i, s in enumerate(p.slides, 1):
        text = "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        for sh in s.shapes:
            if sh.has_table:
                text += "\n" + "\n".join(c.text for r in sh.table.rows for c in r.cells)
        if "？" in text or "?" in text: problems.append(f"slide {i}: question mark on slide")
        if ELABEL.search(text): problems.append(f"slide {i}: evidence label (E0-E3) on slide")
        for b in ban:
            if b in text: problems.append(f"slide {i}: banned phrase '{b}'")
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        m = re.match(r"\s*（约\s*([\d.]+)\s*分钟）", notes)
        if not m: problems.append(f"slide {i}: notes missing '（约 N 分钟）' first line")
        else: total += float(m.group(1))
    print("fonts:", dict(faces))
    print(f"slides: {len(p.slides)}, total minutes: {total:g}")
    print("\n".join(problems) if problems else "all checks passed")

if __name__ == "__main__": main()
